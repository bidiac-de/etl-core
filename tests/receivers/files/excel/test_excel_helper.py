from __future__ import annotations

from pathlib import Path
from typing import Dict, Iterator, List, Optional

import dask.dataframe as dd
import pandas as pd
import pytest
import sys
import types
from typing import Any

import src.etl_core.receivers.files.excel.excel_helper as EH


def _mk_wb_with_rows(fp: Path, rows: List[List[object]], title: str = "Sheet1") -> None:
    """Create a minimal xlsx with given rows using openpyxl."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = title
    for r in rows:
        ws.append(r)
    wb.save(fp)


def test__engine_for_read_and_write_variants_and_errors() -> None:
    assert EH._engine_for_read(".xlsx") == "openpyxl"
    assert EH._engine_for_read(".xlsm") == "openpyxl"

    assert EH._engine_for_read(".xls") == "xlrd"

    # read: unsupported
    with pytest.raises(ValueError):
        EH._engine_for_read(".csv")

    assert EH._engine_for_write(".xlsx") == "openpyxl"
    assert EH._engine_for_write(".xlsm") == "openpyxl"

    with pytest.raises(ValueError):
        EH._engine_for_write(".xls")


def test__read_openpyxl_header_skips_empty_then_uses_first_nonempty() -> None:
    rows_iter = iter(
        [
            tuple(),
            (None, "", None),
            ("A", "B", "C"),
            ("x", "y", "z"),
        ]
    )
    header = EH._read_openpyxl_header(rows_iter)
    assert header == ["A", "B", "C"]

    with pytest.raises(ValueError):
        EH._read_openpyxl_header(iter([tuple(), (None, ""), ()]))


def test__iter_openpyxl_rows_handles_none_row_and_padding(tmp_path: Path) -> None:
    fp = tmp_path / "pad.xlsx"
    _mk_wb_with_rows(
        fp,
        rows=[
            ["col1", "col2", "col3"],
            ["a", "b"],
            ["c", "d", "e"],
        ],
    )

    rows = list(EH._iter_openpyxl_rows(fp, None))

    assert rows[0] == {"col1": "a", "col2": "b", "col3": None}
    assert rows[1] == {"col1": "c", "col2": "d", "col3": "e"}


def test__first_row_empty_all_paths(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    class FakeWS1:
        max_row = 0

        def iter_rows(self, *_, **__):
            return []

    assert EH._first_row_empty(FakeWS1()) is True

    class FakeWS2:
        max_row = 1

        def iter_rows(self, *_, **__):
            return []

    assert EH._first_row_empty(FakeWS2()) is True

    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["a", "b"])
    assert EH._first_row_empty(ws) is False


def test__open_or_create_wb_ws_existing_sheet_and_moves_to_first(
    tmp_path: Path,
) -> None:
    fp = tmp_path / "book.xlsx"
    _mk_wb_with_rows(fp, rows=[["H1"]], title="Sheet1")

    from openpyxl import load_workbook

    wb = load_workbook(fp)
    wb.create_sheet("Another")
    wb.save(fp)
    wb.close()

    wb2, ws2 = EH._open_or_create_wb_ws(fp, "Another")
    try:
        assert ws2.title == "Another"
        assert wb2.worksheets[0].title == "Another"
    finally:
        wb2.close()


def test__open_or_create_wb_ws_move_sheet_fallback(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """
    Force Workbook.move_sheet to raise so the code takes the _sheets fallback path
    (covers lines 182–183).
    """
    fp = tmp_path / "book.xlsx"
    _mk_wb_with_rows(fp, rows=[["H1"]], title="Base")

    from openpyxl.workbook.workbook import Workbook as OB

    def boom(self, *args, **kwargs):  # noqa: D401
        raise RuntimeError("boom")

    monkeypatch.setattr(OB, "move_sheet", boom, raising=True)

    wb3, ws3 = EH._open_or_create_wb_ws(fp, "Target")
    try:
        assert ws3.title == "Target"
        assert wb3.worksheets[0].title == "Target"
    finally:
        wb3.close()


def test_write_excel_row_rejects_non_openpyxl_extension(tmp_path: Path) -> None:
    with pytest.raises(ValueError, match=r"Writing '.xls' is not supported"):
        EH.write_excel_row(tmp_path / "out.xls", {"a": 1})


def test__normalize_to_dataframe_variants() -> None:
    df1 = EH._normalize_to_dataframe([{"a": 1}, {"a": 2}])
    assert (
        isinstance(df1, pd.DataFrame) and list(df1.columns) == ["a"] and len(df1) == 2
    )

    df2 = EH._normalize_to_dataframe([])
    assert isinstance(df2, pd.DataFrame) and df2.empty

    src = pd.DataFrame({"x": [1, 2]})
    df3 = EH._normalize_to_dataframe(src)
    assert df3 is src


def test_read_excel_rows_with_mocked_xlrd_engine(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """
    Simulate the .xls engine without xlrd by monkeypatching
    _prepare_read and _iter_xlrd_rows.
    """
    fake_path = tmp_path / "dummy.xls"

    def fake_prepare(path: Path):
        return fake_path, "xlrd"

    def fake_iter(path: Path, sheet_name: Optional[str]) -> Iterator[Dict[str, object]]:
        yield {"A": 1}
        yield {"A": 2}

    monkeypatch.setattr(EH, "_prepare_read", fake_prepare, raising=True)
    monkeypatch.setattr(EH, "_iter_xlrd_rows", fake_iter, raising=True)

    rows = list(EH.read_excel_rows(fake_path))
    assert rows == [{"A": 1}, {"A": 2}]


def test_read_excel_rows_unknown_engine_raises(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    def fake_prepare(path: Path):
        return tmp_path / "x.abc", "weird"

    monkeypatch.setattr(EH, "_prepare_read", fake_prepare, raising=True)
    with pytest.raises(ValueError, match="Unknown engine"):
        list(EH.read_excel_rows(tmp_path / "x.abc"))


def test_read_excel_bigdata_partitions_capped(tmp_path: Path) -> None:
    fp = tmp_path / "tiny.xlsx"
    _mk_wb_with_rows(fp, rows=[["a"], [1], [2], [3]])

    ddf = EH.read_excel_bigdata(fp, npartitions=99)
    assert isinstance(ddf, dd.DataFrame)
    pdf = ddf.compute()
    assert list(pdf.columns) == ["a"]
    assert pdf["a"].tolist() == [1, 2, 3]


def test__iter_openpyxl_rows_skips_none_and_pads(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """
    Force _iter_openpyxl_rows to see a header, then a short row (pad with None),
    then a None row (skip), then a full row. Covers lines 80 and 83.
    """

    # Fake worksheet iter_rows generator
    class _FakeWS:
        def iter_rows(self, *, values_only: bool = True) -> Iterator:
            assert values_only is True
            # header, short, None, full
            yield ("col1", "col2", "col3")
            yield ("a", "b")
            yield None
            yield ("c", "d", "e")

    class _FakeWB:
        def __init__(self) -> None:
            self._ws = _FakeWS()

        @property
        def worksheets(self) -> List[_FakeWS]:
            return [self._ws]

        def __getitem__(self, _name: str) -> _FakeWS:
            return self._ws

        def close(self) -> None:
            return None

    def _fake_load_workbook(
        _path: Path, *, read_only: bool, data_only: bool
    ) -> _FakeWB:
        assert read_only is True and data_only is True
        return _FakeWB()

    # Inject a tiny "openpyxl" module that only exposes load_workbook()
    fake_openpyxl = types.SimpleNamespace(load_workbook=_fake_load_workbook)
    monkeypatch.setitem(sys.modules, "openpyxl", fake_openpyxl)

    rows = list(EH._iter_openpyxl_rows(Path("dummy.xlsx"), None))
    assert rows == [
        {"col1": "a", "col2": "b", "col3": None},  # padded
        {"col1": "c", "col2": "d", "col3": "e"},  # full
    ]


def test__iter_xlrd_rows_empty_sheet_raises(monkeypatch: pytest.MonkeyPatch) -> None:
    """Covers line 98: 'Worksheet is empty.'."""

    class _EmptySheet:
        nrows = 0
        ncols = 0

    class _Book:
        def sheet_by_index(self, _i: int) -> _EmptySheet:
            return _EmptySheet()

        def sheet_by_name(self, _n: Optional[str]) -> _EmptySheet:
            return _EmptySheet()

    def _open_workbook(_path: Path) -> _Book:
        return _Book()

    fake_xlrd = types.SimpleNamespace(open_workbook=_open_workbook)
    monkeypatch.setitem(sys.modules, "xlrd", fake_xlrd)

    with pytest.raises(ValueError, match="Worksheet is empty."):
        list(EH._iter_xlrd_rows(Path("any.xls"), None))


def test__iter_xlrd_rows_no_header_row_raises(monkeypatch: pytest.MonkeyPatch) -> None:

    class _NoHeaderSheet:
        nrows = 2
        ncols = 2

        def cell_value(self, r: int, c: int) -> Any:
            return "" if r == 0 else f"v{r}{c}"

    class _Book:
        def sheet_by_index(self, _i: int) -> _NoHeaderSheet:
            return _NoHeaderSheet()

        def sheet_by_name(self, _n: Optional[str]) -> _NoHeaderSheet:
            return _NoHeaderSheet()

    fake_xlrd = types.SimpleNamespace(open_workbook=lambda _p: _Book())
    monkeypatch.setitem(sys.modules, "xlrd", fake_xlrd)

    with pytest.raises(ValueError, match="No header row found"):
        list(EH._iter_xlrd_rows(Path("any.xls"), None))


def test__iter_xlrd_rows_yields_rows(monkeypatch: pytest.MonkeyPatch) -> None:
    """Covers the happy-path yield loop (lines 95–107)."""

    class _Sheet:
        nrows = 3
        ncols = 2

        def cell_value(self, r: int, c: int) -> Any:
            # header: "A","B"; rows: (1,2), (3,4)
            table = [
                ["A", "B"],
                [1, 2],
                [3, 4],
            ]
            return table[r][c]

    class _Book:
        def sheet_by_index(self, _i: int) -> _Sheet:
            return _Sheet()

        def sheet_by_name(self, _n: Optional[str]) -> _Sheet:
            return _Sheet()

    fake_xlrd = types.SimpleNamespace(open_workbook=lambda _p: _Book())
    monkeypatch.setitem(sys.modules, "xlrd", fake_xlrd)

    out = list(EH._iter_xlrd_rows(Path("any.xls"), None))
    assert out == [{"A": 1, "B": 2}, {"A": 3, "B": 4}]


def test_write_excel_row_engine_guard_branch(
    monkeypatch: pytest.MonkeyPatch, tmp_path: Path
) -> None:
    """
    Hit write_excel_row's own engine guard (lines 201–203) by bypassing the
    earlier _engine_for_write ValueError.
    """
    fake_path = tmp_path / "out.xlsx"

    def _fake_prepare_write(p: Path) -> Any:
        # Return a non-openpyxl engine so the guard raises from inside write_excel_row
        return fake_path, "xlrd"

    monkeypatch.setattr(EH, "_prepare_write", _fake_prepare_write)
    with pytest.raises(ValueError, match=r"only supported"):
        EH.write_excel_row(fake_path, {"a": 1})
